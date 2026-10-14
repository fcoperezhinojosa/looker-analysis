"""Diff two capture workbooks cell by cell.

Both files come out of the same generator, so any difference in the
hard-coded cells is a difference in what the capture read out of Looker.
That is the point: re-capturing a period we already have tells us whether
the pipeline is repeatable, and if it is not, which tile moved.

Only non-formula cells are compared. Rates and deltas are derived, so
diffing them would just re-report the same discrepancy several times over,
and their cached values depend on whether the file was ever opened in Excel.

  python scripts/compare_captures.py REFERENCE.xlsx FRESH.xlsx
"""
import sys
from pathlib import Path

import openpyxl

SHEETS = ("Totals", "Countries")


def cells(path):
    """(sheet, row, col) -> value, for hard-coded cells only."""
    formulas = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    out, labels = {}, {}
    for name in SHEETS:
        if name not in formulas.sheetnames:
            continue
        fs, vs = formulas[name], values[name]
        for row in fs.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    continue
                v = vs.cell(row=c.row, column=c.column).value
                if isinstance(v, (int, float)):
                    out[(name, c.row, c.column)] = v
        # row label = the leading text columns, for readable output
        for row in vs.iter_rows(values_only=False):
            txt = [c.value for c in row[:4] if isinstance(c.value, str)]
            if txt:
                labels[(name, row[0].row)] = " / ".join(txt)
    return out, labels


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    ref_p, new_p = Path(sys.argv[1]), Path(sys.argv[2])
    ref, labels = cells(ref_p)
    new, _ = cells(new_p)

    same = diff = 0
    only_ref = sorted(set(ref) - set(new))
    only_new = sorted(set(new) - set(ref))
    worst = []

    for k in sorted(set(ref) & set(new)):
        a, b = ref[k], new[k]
        if a == b:
            same += 1
            continue
        diff += 1
        pct = (b / a - 1) * 100 if a else float("inf")
        worst.append((abs(pct), k, a, b, pct))

    print(f"reference : {ref_p.name}")
    print(f"fresh     : {new_p.name}")
    print(f"\nhard-coded cells compared : {same + diff}")
    print(f"  identical               : {same}")
    print(f"  different               : {diff}")
    if only_ref:
        print(f"  present only in reference: {len(only_ref)}")
    if only_new:
        print(f"  present only in fresh    : {len(only_new)}")

    if diff:
        print(f"\nthe {min(30, len(worst))} largest differences:")
        worst.sort(reverse=True)
        for _, k, a, b, pct in worst[:30]:
            sheet, row, col = k
            lab = labels.get((sheet, row), "")
            print(f"  {sheet:9s} r{row:<4d} c{col:<3d} {a:>10,} -> {b:>10,} "
                  f"({pct:+7.1f}%)  {lab}")
    else:
        print("\nno differences: the capture reproduced the reference exactly.")

    for tag, keys in (("reference", only_ref), ("fresh", only_new)):
        if keys:
            print(f"\ncells only in {tag} (a blank in the other file):")
            for k in keys[:20]:
                sheet, row, col = k
                print(f"  {sheet:9s} r{row:<4d} c{col:<3d}  {labels.get((sheet,row),'')}")

    return 1 if diff or only_ref or only_new else 0


if __name__ == "__main__":
    sys.exit(main())
